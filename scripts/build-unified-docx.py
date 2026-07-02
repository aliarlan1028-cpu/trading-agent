from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "AI_Trading_Agent_统一需求文档.md"
OUT = ROOT / "docs" / "AI_Trading_Agent_统一需求文档.docx"


def set_cell(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    for margin_name in ["top", "bottom", "start", "end"]:
        node = OxmlElement(f"w:{margin_name}Mar")
        node.set(qn("w:w"), "120" if margin_name in ["start", "end"] else "80")
        node.set(qn("w:type"), "dxa")
        tc_pr.append(node)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_geometry(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.tblW
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    tbl_pr.append(tbl_ind)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def flush_table(doc, rows):
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    set_table_geometry(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            value = row[c_idx] if c_idx < len(row) else ""
            set_cell(table.cell(r_idx, c_idx), value, bold=r_idx == 0)
            if r_idx == 0:
                shade_cell(table.cell(r_idx, c_idx), "F2F4F7")


def parse_markdown(doc, text):
    table_rows = []
    in_code = False
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            continue
        if not line.strip():
            flush_table(doc, table_rows)
            table_rows = []
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
                table_rows.append(cells)
            continue
        flush_table(doc, table_rows)
        table_rows = []
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(24)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("0B2545")
            p.paragraph_format.space_after = Pt(10)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
        elif re.match(r"^\d+\.\s+", line):
            add_number(doc, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            add_bullet(doc, line[2:].strip())
        else:
            doc.add_paragraph(line)
    flush_table(doc, table_rows)


def main():
    doc = Document()
    configure_styles(doc)
    parse_markdown(doc, MD.read_text(encoding="utf-8"))
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("AI Trading Agent 统一需求文档").font.size = Pt(9)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
